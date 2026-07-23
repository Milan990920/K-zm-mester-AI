import csv
import hashlib
import io
from datetime import date

from openpyxl import load_workbook

from app.models.document import Document
from app.models.enums import DeliveryFormat, DocumentStatus, InvoiceType, UtilityType
from app.models.invoice import Invoice


def _make_invoice(db_session, tenant, user, *, invoice_number: str, gross_amount: float) -> Invoice:
    content = f"{tenant.id}-{invoice_number}".encode()
    document = Document(
        tenant_id=tenant.id,
        uploaded_by_user_id=user.id,
        original_filename="teszt.pdf",
        gcs_path="tenants/x/documents/y.pdf",
        content_hash=hashlib.sha256(content).hexdigest(),
        file_size_bytes=10,
        status=DocumentStatus.DONE,
    )
    db_session.add(document)
    db_session.flush()

    invoice = Invoice(
        tenant_id=tenant.id,
        document_id=document.id,
        utility_type=UtilityType.ELECTRICITY,
        invoice_type=InvoiceType.COMMERCIAL,
        delivery_format=DeliveryFormat.ELECTRONIC,
        invoice_number=invoice_number,
        invoice_date=date(2026, 1, 15),
        gross_amount=gross_amount,
        currency="HUF",
    )
    db_session.add(invoice)
    db_session.commit()
    return invoice


def _login(client, email, password):
    response = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    return response.json()["access_token"]


def test_export_csv_contains_invoice_rows(client, make_tenant, make_user, db_session):
    tenant = make_tenant(name="Export Kft.")
    user = make_user(tenant.id, "customer_admin", email="export1@teszt.hu", password="Titok1234!")
    token = _login(client, "export1@teszt.hu", "Titok1234!")

    _make_invoice(db_session, tenant, user, invoice_number="EXP-001", gross_amount=1000.0)
    _make_invoice(db_session, tenant, user, invoice_number="EXP-002", gross_amount=2000.0)

    response = client.get(
        "/api/v1/invoices/export?format=csv", headers={"Authorization": f"Bearer {token}"}
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]

    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0][0] == "Számlaszám"
    invoice_numbers = [row[0] for row in rows[1:]]
    assert set(invoice_numbers) == {"EXP-001", "EXP-002"}


def test_export_xlsx_contains_invoice_rows(client, make_tenant, make_user, db_session):
    tenant = make_tenant(name="Export XLSX Kft.")
    user = make_user(tenant.id, "customer_admin", email="export2@teszt.hu", password="Titok1234!")
    token = _login(client, "export2@teszt.hu", "Titok1234!")

    _make_invoice(db_session, tenant, user, invoice_number="XLSX-001", gross_amount=1500.0)

    response = client.get(
        "/api/v1/invoices/export?format=xlsx", headers={"Authorization": f"Bearer {token}"}
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "Számlaszám"
    assert rows[1][0] == "XLSX-001"
    # Enum columns must be plain values ("electricity"), never the Python
    # repr ("UtilityType.ELECTRICITY") that openpyxl's str() coercion gives
    # a `class X(str, Enum)` member by default.
    assert rows[1][1] == "electricity"
    assert rows[1][2] == "commercial"


def test_export_only_includes_own_tenant(client, make_tenant, make_user, db_session):
    tenant_a = make_tenant(name="Export A Kft.")
    tenant_b = make_tenant(name="Export B Kft.")
    user_a = make_user(tenant_a.id, "customer_admin", email="export-a@teszt.hu", password="Titok1234!")
    user_b = make_user(tenant_b.id, "customer_admin", email="export-b@teszt.hu", password="Titok1234!")

    _make_invoice(db_session, tenant_a, user_a, invoice_number="A-001", gross_amount=100.0)
    _make_invoice(db_session, tenant_b, user_b, invoice_number="B-001", gross_amount=200.0)

    token_a = _login(client, "export-a@teszt.hu", "Titok1234!")
    response = client.get(
        "/api/v1/invoices/export?format=csv", headers={"Authorization": f"Bearer {token_a}"}
    )

    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    invoice_numbers = [row[0] for row in rows[1:]]
    assert invoice_numbers == ["A-001"]
